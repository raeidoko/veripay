from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "VeriPay_Investor_Ask_and_Business_Model.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
GREEN_FILL = "EAF7EF"
GOLD_FILL = "FFF7E0"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(table, top=80, start=120, bottom=80, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.find(qn("w:tblCellMar"))
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for key, value in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = margins.find(qn(f"w:{key}"))
        if node is None:
            node = OxmlElement(f"w:{key}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths, indent=120):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(indent))
    tbl_ind.set(qn("w:type"), "dxa")

    old_grid = tbl.tblGrid
    if old_grid is not None:
        tbl.remove(old_grid)
    grid = OxmlElement("w:tblGrid")
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    tbl.insert(1, grid)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            width = widths[min(idx, len(widths) - 1)]
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.width = Inches(width / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP
    set_cell_margins(table)


def style_cell_text(cell, bold=False, color=NAVY, size=10):
    for paragraph in cell.paragraphs:
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.1
        for run in paragraph.runs:
            set_run_font(run, size=size, color=color, bold=bold)


def add_para(doc, text="", style=None, bold_start=None):
    paragraph = doc.add_paragraph(style=style)
    paragraph.paragraph_format.space_after = Pt(6)
    paragraph.paragraph_format.line_spacing = 1.1
    if bold_start and text.startswith(bold_start):
        run = paragraph.add_run(bold_start)
        set_run_font(run, size=11, color=NAVY, bold=True)
        rest = paragraph.add_run(text[len(bold_start):])
        set_run_font(rest, size=11, color=NAVY)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11, color=NAVY)
    return paragraph


def add_bullet(doc, text):
    paragraph = doc.add_paragraph(style="List Bullet")
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    return paragraph


def add_numbered(doc, text):
    paragraph = doc.add_paragraph(style="List Number")
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.line_spacing = 1.167
    run = paragraph.add_run(text)
    set_run_font(run, size=11, color=NAVY)
    return paragraph


def add_callout(doc, label, text, fill=CALLOUT_FILL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, fill)
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    label_run = paragraph.add_run(label + ": ")
    set_run_font(label_run, size=11, color=DARK_BLUE, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color=NAVY)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    set_table_geometry(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = header
        set_cell_shading(cell, LIGHT_FILL)
        style_cell_text(cell, bold=True, color=NAVY, size=10)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
            style_cell_text(cells[idx], color=NAVY, size=10)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return table


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.1

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("VeriPay Investor Brief")
    set_run_font(run, size=9, color=MUTED)


def build():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("VeriPay Investor Ask and Business Model")
    set_run_font(run, size=20, color=NAVY, bold=True)

    subtitle = doc.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("A practical funding and revenue strategy for an early escrow payment platform")
    set_run_font(run, size=12, color=MUTED)

    add_callout(
        doc,
        "Core recommendation",
        "Position VeriPay for a pre-seed raise of about $150,000 to $300,000, unless stronger traction supports a larger round. The business model should start with free buyer access, seller verification revenue, and a 2% to 3% escrow service fee on completed transactions.",
        GREEN_FILL,
    )

    doc.add_heading("1. Executive Summary", level=1)
    add_para(
        doc,
        "VeriPay is best presented to investors as a trust infrastructure product for online and social-commerce transactions. The product protects buyers by holding funds until agreed delivery conditions are met, while helping sellers prove legitimacy, reduce buyer hesitation, and create more professional transaction records.",
    )
    add_para(
        doc,
        "At this stage, the investor conversation should focus less on asking for the biggest possible cheque and more on asking for the right amount to reach specific proof points. Investors want to know what the capital will unlock, how the business will make money, and what evidence will show that VeriPay can scale.",
    )

    doc.add_heading("2. Recommended Investor Ask", level=1)
    add_para(
        doc,
        "The strongest way to ask investors is to connect the funding amount to milestones. A good framing is: Based on our current stage, product scope, and market opportunity, what pre-seed funding amount would be appropriate, and what milestones should that money help us reach?",
    )
    add_table(
        doc,
        ["Funding range", "When it makes sense", "What it should help build"],
        [
            (
                "$50,000 to $150,000",
                "Useful for a small angel or friends-and-family round when the product is still being validated.",
                "Production backend, basic payment flow research, security review, early seller onboarding, and pilot testing.",
            ),
            (
                "$150,000 to $300,000",
                "Recommended starting ask for VeriPay if there is a working product, clear market pain, and early merchant interest.",
                "Backend/API, payment-provider integration, seller verification operations, dispute process design, legal/compliance preparation, and go-to-market testing.",
            ),
            (
                "$300,000 to $500,000",
                "Appropriate if VeriPay can show stronger traction such as active merchants, pilot users, partnerships, or transaction volume.",
                "Small team hiring, full production launch, compliance work, merchant acquisition, support operations, and measurable growth experiments.",
            ),
        ],
        [1800, 3300, 4260],
    )

    add_callout(
        doc,
        "Recommended ask",
        "For the current stage, VeriPay should likely ask for $150,000 to $300,000 as a pre-seed round. A larger ask can work if there is clear evidence of traction, signed pilot partners, or strong market demand.",
        GOLD_FILL,
    )

    doc.add_heading("3. How to Explain the Ask to Investors", level=1)
    add_para(doc, "A clean investor explanation should sound specific, disciplined, and milestone-based:")
    add_numbered(doc, "We are raising $150,000 to $300,000 to move VeriPay from a front-end product into a secure pilot-ready escrow platform.")
    add_numbered(doc, "The capital will fund payment integration work, backend development, seller verification, dispute operations, early compliance, and merchant acquisition.")
    add_numbered(doc, "The goal is to prove that buyers and sellers will use escrow for high-trust online transactions and that VeriPay can earn predictable fee revenue from completed transactions.")
    add_numbered(doc, "Success will be measured by onboarded sellers, completed escrow transactions, transaction value, repeat usage, dispute-resolution speed, and revenue per transaction.")

    doc.add_heading("4. Best Business Model", level=1)
    add_para(
        doc,
        "The right business model for VeriPay should be simple enough for users to understand and strong enough for investors to see scalability. The best starting model is a blended transaction and merchant-services model.",
    )
    add_callout(
        doc,
        "Best starting model",
        "Free buyer use plus seller verification revenue plus a 2% to 3% escrow transaction fee. This aligns revenue with successful transactions and keeps the buyer experience friction-light.",
        GREEN_FILL,
    )

    doc.add_heading("5. Revenue Streams", level=1)
    add_table(
        doc,
        ["Revenue stream", "How it works", "Why it fits VeriPay"],
        [
            (
                "Escrow transaction fee",
                "VeriPay charges a small percentage, ideally 2% to 3%, on successful completed transactions.",
                "This is the core model because VeriPay earns when it successfully creates trust and completes a transaction.",
            ),
            (
                "Seller verification fee",
                "Sellers pay a monthly, yearly, or one-time fee to become verified and display trust credentials.",
                "Trust is the main value proposition, so sellers have a reason to pay for credibility.",
            ),
            (
                "Premium seller subscription",
                "Active sellers pay for higher limits, better receipt tools, analytics, dispute priority, and profile visibility.",
                "This turns serious merchants into recurring revenue customers.",
            ),
            (
                "Dispute resolution fee",
                "A small admin fee can apply to complex mediated disputes, especially when extra review is required.",
                "It helps cover operational cost, but it must be used carefully so customers do not feel punished for reporting issues.",
            ),
            (
                "Logistics partner commission",
                "VeriPay can earn a referral or booking commission from delivery partners used inside the checkout flow.",
                "Delivery is closely tied to escrow release, so logistics partnerships can add revenue and improve completion confidence.",
            ),
        ],
        [2100, 3650, 3610],
    )

    doc.add_heading("6. Example Pricing Structure", level=1)
    add_table(
        doc,
        ["Plan or fee", "Suggested pricing", "Purpose"],
        [
            ("Buyer account", "Free", "Keep buyer adoption easy and remove friction from starting transactions."),
            ("Escrow service fee", "2% to 3% per completed transaction", "Primary revenue stream tied directly to transaction value."),
            ("Verified seller", "Monthly or yearly verification fee", "Monetize trust and legitimacy for sellers."),
            ("Pro seller subscription", "Tiered monthly fee", "Offer transaction limits, receipts, analytics, and priority support."),
            ("Business seller plan", "Higher monthly fee for teams or high-volume sellers", "Support serious merchants and future enterprise-style features."),
        ],
        [2200, 2900, 4260],
    )

    doc.add_heading("7. Use of Funds", level=1)
    add_para(doc, "Investors will expect the funding ask to map to clear spending categories. VeriPay can explain the raise around these priorities:")
    add_bullet(doc, "Product engineering: Build the backend, API, database, authentication, admin controls, and production-ready transaction records.")
    add_bullet(doc, "Payment and escrow integration: Evaluate licensed providers, settlement flows, wallet behavior, refunds, and audit trails.")
    add_bullet(doc, "Security and compliance: Strengthen identity checks, fraud controls, dispute records, data protection, and legal review.")
    add_bullet(doc, "Seller onboarding: Recruit early verified sellers in high-risk categories such as gadgets, fashion, beauty products, services, and marketplace sales.")
    add_bullet(doc, "Customer support and dispute operations: Define response times, mediation rules, evidence handling, and refund/release processes.")
    add_bullet(doc, "Go-to-market testing: Run small campaigns, measure conversion, and learn which customer segment adopts escrow fastest.")

    doc.add_heading("8. Milestones Investors Should Expect", level=1)
    add_table(
        doc,
        ["Milestone", "What good progress looks like"],
        [
            ("Pilot launch", "A working escrow flow with selected buyers and verified sellers using the product in real transactions."),
            ("Merchant traction", "A growing number of sellers complete verification and actively invite buyers to use VeriPay."),
            ("Transaction volume", "Clear movement in completed escrow transactions and gross transaction value."),
            ("Repeat usage", "Buyers and sellers come back because the trust layer reduces risk and friction."),
            ("Revenue proof", "Transaction fees and seller payments show that the model can generate repeatable revenue."),
            ("Operational confidence", "Disputes are resolved clearly, fairly, and quickly enough to protect trust in the platform."),
        ],
        [2800, 6560],
    )

    doc.add_heading("9. Investor-Friendly Pitch Language", level=1)
    add_para(
        doc,
        "VeriPay makes money by charging a small escrow service fee on completed transactions, with additional revenue from seller verification, premium merchant tools, and logistics partnerships.",
    )
    add_para(
        doc,
        "The platform starts with the trust problem in informal and online commerce: buyers fear losing money, and legitimate sellers struggle to prove they can be trusted. VeriPay solves this by holding funds safely, documenting the agreement, guiding delivery confirmation, and creating a process for disputes when something goes wrong.",
    )
    add_para(
        doc,
        "The funding ask should show discipline: VeriPay is not raising money just to build more features. It is raising capital to prove a specific commercial thesis: that escrow can increase transaction confidence, seller conversion, and repeat online trade while producing fee-based revenue.",
    )

    doc.add_heading("10. Risks and How to Address Them", level=1)
    add_table(
        doc,
        ["Risk", "Investor concern", "How VeriPay should respond"],
        [
            ("Regulatory and payment compliance", "Escrow and payments can be sensitive areas.", "Work with licensed payment partners, legal advisors, and clear transaction policies before live money movement."),
            ("Fraud and bad actors", "A trust product can attract misuse.", "Use seller verification, buyer/seller history, evidence records, limits, review workflows, and suspicious-activity flags."),
            ("Dispute operations", "Manual disputes can become expensive.", "Start with clear rules and templates, then automate common decisions as transaction volume grows."),
            ("Adoption friction", "Users may resist paying fees.", "Keep buyers free, make fees transparent, and show that escrow reduces the risk of losing the full purchase amount."),
        ],
        [2050, 3100, 4210],
    )

    doc.add_heading("11. Final Recommendation", level=1)
    add_para(
        doc,
        "VeriPay should approach investors with a $150,000 to $300,000 pre-seed ask, framed around specific milestones rather than general product ambition. The ask should be supported by a clear business model: free buyers, verified sellers, a 2% to 3% escrow service fee, seller subscriptions, and potential logistics commissions.",
    )
    add_para(
        doc,
        "This model is easy to explain, aligned with the trust problem VeriPay solves, and scalable if transaction volume grows. The most important next step is proving that real buyers and sellers will complete transactions through VeriPay and that sellers are willing to pay for verification and trust-building tools.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
