from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "VeriPay_Technology_Stack.docx"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin_name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin_name}"))
        if node is None:
            node = OxmlElement(f"w:{margin_name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.right_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("VeriPay Technology Stack")
    set_run_font(r, size=9, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.text = ""
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = footer.add_run("Prepared for VeriPay project documentation")
    set_run_font(r, size=8.5, color=MUTED)


def add_para(doc, text="", size=11, bold=False, italic=False, color=None, after=6, before=0, align=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    r = p.add_run(text)
    set_run_font(r, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11)


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    r = p.add_run(text)
    set_run_font(r, size=11)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    table.style = "Table Grid"
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(header)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            p = cells[idx].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            r = p.add_run(value)
            set_run_font(r, size=10)
    set_table_widths(table, widths)
    add_para(doc, "", after=4)


def build():
    doc = Document()
    style_doc(doc)

    add_para(doc, "VERIPAY TECHNICAL DOCUMENTATION", size=10, bold=True, color=BLUE, after=2)
    add_para(doc, "Technology Stack and Usage Report", size=24, bold=True, color=NAVY, after=4)
    add_para(
        doc,
        "A detailed explanation of every major technology used to build the VeriPay escrow payment platform and what each technology contributes to the product.",
        size=13,
        color=MUTED,
        after=16,
    )

    for label, value in [
        ("Project", "VeriPay"),
        ("Application type", "Single-page React web application"),
        ("Primary language", "TypeScript with TSX components"),
        ("Build system", "Vite"),
        ("UI styling", "Tailwind CSS v4 with custom CSS theme tokens"),
        ("Document date", "August 14, 2026"),
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2)

    add_callout(
        doc,
        "Short summary",
        "VeriPay is built as a modern browser-based escrow dashboard. React provides the component structure, TypeScript defines the data contracts, Vite handles the local development and production build, Tailwind CSS provides the visual system, and browser APIs provide client-side storage, sound, date/currency formatting, printing, and basic interaction behavior.",
    )

    doc.add_heading("1. High-Level Architecture", level=1)
    add_para(
        doc,
        "The project is a client-side single-page application. The browser loads index.html, Vite mounts the React application through src/main.tsx, and the main App component coordinates transaction state, seller selection, notification state, dispute handling, receipt display, and routing between major dashboard views.",
    )
    add_para(
        doc,
        "There is no dedicated backend server in the current project. Transaction records, alerts, and user interactions are handled in browser memory and persisted through localStorage. Seed data is stored in TypeScript files so the product can show realistic escrow records, seller profiles, and dispute cases immediately after loading.",
    )

    add_table(
        doc,
        ["Layer", "Main technology", "What it builds in VeriPay"],
        [
            ("Application shell", "React, React DOM, Vite", "Loads the dashboard, mounts the app, and renders the single-page interface."),
            ("Business state", "React hooks and TypeScript types", "Stores transactions, notifications, selected seller, active tab, receipt modal, and dispute decisions."),
            ("User interface", "TSX components, Tailwind CSS, Lucide React", "Builds the escrow pipeline, seller directory, dispute center, alerts, receipt modal, forms, icons, and responsive layout."),
            ("Data model", "TypeScript interfaces and seed data", "Defines transaction statuses, evidence files, sellers, disputes, notifications, and timeline events."),
            ("Browser services", "localStorage, Web Audio API, Intl APIs, print APIs", "Persists records locally, plays feedback sounds, formats currency/dates, and prints receipts."),
            ("Build and validation", "npm scripts, TypeScript compiler, Vite build", "Runs local development server, validates types, and creates the production dist bundle."),
        ],
        [2000, 2600, 4760],
    )

    doc.add_heading("2. Runtime and Application Framework", level=1)
    add_table(
        doc,
        ["Technology", "Installed version", "Where it appears", "What it was used to build"],
        [
            ("React", "19.2.7", "src/App.tsx and all component files", "The main UI framework. It powers component rendering, state updates, event handlers, modals, forms, conditional screens, and reusable dashboard sections."),
            ("React DOM", "19.2.7", "src/main.tsx", "Mounts the React application into the root DOM element in index.html."),
            ("JSX / TSX", "Part of React + TypeScript workflow", "src/*.tsx and src/components/*.tsx", "Allows UI markup and TypeScript logic to live together inside component files."),
            ("React StrictMode", "React built-in", "src/main.tsx", "Wraps the app during rendering to help expose unsafe component patterns during development."),
        ],
        [1700, 1500, 2500, 3660],
    )

    add_para(
        doc,
        "React is the central technology behind the product experience. Each major page is implemented as a component: EscrowWizard handles agreement creation and active escrow workflows, SellerDirectory handles merchant discovery, DisputeCenter handles claims and arbitration, NotificationsPanel handles alerts, DigitalReceipt handles settlement proof, and VeriPayLogo handles brand identity.",
    )

    doc.add_heading("3. TypeScript and Data Modeling", level=1)
    add_para(
        doc,
        "TypeScript is used to make the application safer and easier to maintain. The project defines clear interfaces for every major business object in src/types.ts. These interfaces help prevent mismatched data shapes when the UI passes transactions, sellers, evidence files, disputes, and notifications between components.",
    )
    add_table(
        doc,
        ["TypeScript model", "Purpose in the app"],
        [
            ("TransactionStatus", "Defines the allowed transaction lifecycle states: draft, pending payment, secured, dispatched, in transit, delivered, completed, disputed, and refunded."),
            ("TimelineEvent", "Stores each event in the transaction history, including status, title, description, timestamp, and actor."),
            ("SellerProfile", "Models each verified seller with trust score, rating, phone number, category, active escrow count, and sales history."),
            ("Transaction", "Represents an escrow order, including parties, product, amount, logistics partner, status, timeline, dispute, and receipt-related fields."),
            ("Dispute", "Stores claim reason, description, buyer evidence, seller evidence, messages, arbitration status, and final verdict description."),
            ("Notification", "Stores dashboard alerts with title, message, transaction ID, type, timestamp, and read/unread state."),
        ],
        [2400, 6960],
    )

    doc.add_heading("4. Build Tooling and Developer Workflow", level=1)
    add_table(
        doc,
        ["Technology", "Installed version", "Use in the project"],
        [
            ("Vite", "6.4.3", "Provides the local development server, hot module behavior, module bundling, and production build output in dist/."),
            ("@vitejs/plugin-react", "5.2.0", "Enables React support inside Vite, including TSX transformation and React development behavior."),
            ("@tailwindcss/vite", "4.3.2", "Connects Tailwind CSS v4 directly into the Vite build pipeline."),
            ("TypeScript", "5.8.3", "Runs type checking through npm run lint, using tsc --noEmit."),
            ("esbuild", "0.25.12", "Used by Vite internally for fast TypeScript/JavaScript transformation and dependency optimization."),
            ("npm", "Project package manager", "Installs dependencies, stores package-lock.json, and runs dev, build, preview, clean, and lint scripts."),
            ("Node.js", "Runtime required by tooling", "Executes npm scripts, Vite, TypeScript, and package tooling."),
        ],
        [2300, 1500, 5560],
    )
    add_para(doc, "The project exposes the following npm scripts:")
    for script in [
        "npm run dev starts Vite on port 3000 and exposes the app on the local network with host 0.0.0.0.",
        "npm run build creates the production-ready static bundle inside dist/.",
        "npm run preview serves the built output for local production-style testing.",
        "npm run lint runs TypeScript type checking without emitting files.",
        "npm run clean removes generated build artifacts.",
    ]:
        add_bullet(doc, script)

    doc.add_heading("5. Styling and Visual Design Technologies", level=1)
    add_table(
        doc,
        ["Technology", "Installed version", "What it was used to build"],
        [
            ("Tailwind CSS", "4.3.2", "Provides utility classes for layout, spacing, typography, color, borders, shadows, responsive grids, buttons, cards, badges, and modal surfaces."),
            ("Tailwind @theme", "Tailwind CSS v4 feature", "Defines the app font tokens for Inter and JetBrains Mono in src/index.css."),
            ("Custom CSS", "Browser CSS", "Defines imported fonts, body background, global smoothing, and scrollbar styling."),
            ("Google Fonts", "Inter and JetBrains Mono", "Provides the main interface typeface and monospaced styling for IDs, references, and technical-looking labels."),
            ("Responsive CSS utilities", "Tailwind responsive classes", "Builds mobile-first layouts using grid, flexbox, sm/md/lg breakpoints, overflow handling, and adaptive controls."),
        ],
        [2400, 1900, 5060],
    )
    add_para(
        doc,
        "The visual system uses a restrained fintech palette: deep navy for trust and authority, blue for primary actions, emerald for successful escrow/payment states, amber for pending or warning states, rose for dispute/risk states, and slate neutrals for structure and readability.",
    )

    doc.add_heading("6. Icons, SVG, and Visual Assets", level=1)
    add_table(
        doc,
        ["Technology or asset type", "Installed version / source", "What it was used to build"],
        [
            ("Lucide React", "0.546.0", "Provides consistent icon components for navigation tabs, seller actions, receipts, notifications, dispute center, wallet actions, delivery tracking, and status indicators."),
            ("Inline SVG", "Native browser SVG", "Builds the VeriPay shield logo and generated QR-style payment graphic without requiring external image files."),
            ("Remote image URLs", "Unsplash image URLs in seed data", "Represent buyer and seller dispute evidence attachments in the arbitration center."),
            ("CSS shapes and badges", "Tailwind and CSS", "Creates status pills, trust score badges, stepper dots, alert indicators, card borders, modal backdrops, and section emphasis."),
        ],
        [2600, 2000, 4760],
    )

    doc.add_heading("7. Browser APIs Used in the Product", level=1)
    add_table(
        doc,
        ["Browser API", "Where it is used", "Purpose"],
        [
            ("localStorage", "src/App.tsx", "Persists transactions and notifications between page reloads. It also recovers gracefully from invalid stored data."),
            ("Web Audio API", "src/App.tsx", "Creates short success, alert, and click sounds using AudioContext, OscillatorNode, and GainNode."),
            ("Intl.NumberFormat", "src/components/DigitalReceipt.tsx", "Formats transaction value, escrow fee, and total settled amount as Nigerian naira currency."),
            ("Date and time formatting", "App, receipt, history, dispute center", "Formats created dates, settlement dates, notification timestamps, dispute filing time, and chat timestamps."),
            ("window.open and print", "src/components/DigitalReceipt.tsx", "Creates a printable receipt window for the completed transaction document."),
            ("DOM root mounting", "src/main.tsx and index.html", "Mounts the React application into the root element."),
            ("setTimeout", "App and DisputeCenter", "Delays receipt opening after settlement and adds a mediator response after a case message."),
        ],
        [2300, 2700, 4360],
    )

    doc.add_heading("8. Component-Level Technology Usage", level=1)
    add_table(
        doc,
        ["Component/file", "Technologies used", "What it builds"],
        [
            ("src/App.tsx", "React state/effects, TypeScript, localStorage, Web Audio API, Lucide icons", "The main dashboard shell, app state, transaction lifecycle actions, notifications, tab navigation, dispute resolution, and receipt modal control."),
            ("src/components/EscrowWizard.tsx", "React forms, TypeScript props, Tailwind, SVG, Lucide icons", "Escrow creation form, payment checkout modal, active status timeline, logistics map, delivery confirmation, and dispute form."),
            ("src/components/SellerDirectory.tsx", "React state, filtering logic, Tailwind cards, Lucide icons", "Searchable and filterable verified merchant directory with trust scores and escrow-start actions."),
            ("src/components/DisputeCenter.tsx", "React state/effects, typed transactions, Tailwind grids, image rendering", "Active dispute list, evidence display, case details, mediated discussion, and arbitration decisions."),
            ("src/components/DigitalReceipt.tsx", "React refs, Intl currency formatting, print window, Tailwind modal", "Printable escrow settlement receipt with buyer/seller, amount, fee, tracking reference, terms, and verification badge."),
            ("src/components/NotificationsPanel.tsx", "React props, Tailwind alert cards, Lucide icons", "Escrow alert feed with read/unread handling and status-specific visual styling."),
            ("src/components/VeriPayLogo.tsx", "React props, inline SVG, Tailwind typography", "Reusable brand logo in multiple sizes with optional subtitle."),
            ("src/seedData.ts", "TypeScript constants", "Initial transaction, seller, timeline, dispute, evidence, and status data used to populate the product."),
            ("src/types.ts", "TypeScript interfaces and union types", "Shared data contracts for the entire application."),
        ],
        [2600, 2700, 4060],
    )

    doc.add_heading("9. Data Persistence and Current Storage Approach", level=1)
    add_para(
        doc,
        "The current project stores operational records in the browser rather than a server database. This is handled through localStorage keys for transactions and notifications. The app initializes from localStorage when saved data exists, otherwise it falls back to seed data. It also removes invalid saved data if JSON parsing fails.",
    )
    add_para(
        doc,
        "This approach keeps the app lightweight and easy to run locally. For a full production deployment, the same TypeScript models could map naturally to backend tables or API resources for users, sellers, transactions, disputes, evidence files, notifications, and receipts.",
    )

    doc.add_heading("10. Files and Project Structure", level=1)
    add_table(
        doc,
        ["Path", "Purpose"],
        [
            ("index.html", "HTML entry point with metadata, viewport settings, title, and root div."),
            ("src/main.tsx", "React DOM entry point that mounts App into the page."),
            ("src/App.tsx", "Main app coordinator and top-level dashboard state."),
            ("src/index.css", "Global Tailwind import, font imports, theme font tokens, body styling, and scrollbar styles."),
            ("src/components/", "Reusable UI components for each major product surface."),
            ("src/types.ts", "Application-wide TypeScript data contracts."),
            ("src/seedData.ts", "Initial seller, transaction, timeline, and dispute records."),
            ("vite.config.ts", "Vite configuration for React, Tailwind, aliasing, HMR, and file watching behavior."),
            ("package.json", "Project scripts, runtime dependencies, and development dependencies."),
            ("package-lock.json", "Exact installed dependency tree."),
            ("dist/", "Generated production build output."),
            ("docs/", "Project documentation and generated Word documents."),
        ],
        [2500, 6860],
    )

    doc.add_heading("11. Why These Technologies Fit VeriPay", level=1)
    for reason in [
        "React is well suited for a dashboard with many conditional screens, forms, modals, and reusable cards.",
        "TypeScript is valuable because the escrow domain depends on clear states, roles, disputes, evidence records, and transaction objects.",
        "Vite keeps development and build performance fast, which is useful for frequent UI iteration.",
        "Tailwind CSS supports a consistent visual system without requiring a large custom CSS codebase.",
        "Lucide React keeps iconography professional and consistent across payment, trust, logistics, and dispute screens.",
        "Browser APIs keep the current app lightweight while still allowing persistence, printing, sound feedback, and localized formatting.",
    ]:
        add_bullet(doc, reason)

    doc.add_heading("12. Recommended Future Technologies", level=1)
    add_table(
        doc,
        ["Future need", "Recommended technology category", "Reason"],
        [
            ("Real user accounts", "Authentication provider or custom auth service", "Separate buyer, seller, logistics, and arbitrator permissions."),
            ("Persistent transactions", "Backend API and database", "Move escrow records from browser localStorage to secure server-side storage."),
            ("Payment execution", "Licensed payment provider integration", "Fund escrow, release seller payouts, refund buyers, and audit settlement events."),
            ("Evidence uploads", "Object storage and file-scanning workflow", "Store dispute photos, documents, videos, and courier proof securely."),
            ("Realtime updates", "WebSockets or realtime database subscriptions", "Push transaction, delivery, payment, and dispute changes to all parties."),
            ("Testing", "Unit, component, and end-to-end test tooling", "Protect payment and dispute workflows from regressions."),
            ("Deployment", "Static hosting or full-stack app hosting", "Serve the Vite build reliably with HTTPS and environment-specific configuration."),
        ],
        [2300, 3100, 3960],
    )

    doc.add_heading("Conclusion", level=1)
    add_para(
        doc,
        "VeriPay uses a focused modern front-end stack: React for the interface, TypeScript for data safety, Vite for development and build tooling, Tailwind CSS for styling, Lucide React for iconography, and browser APIs for local persistence and receipt behavior. Together, these technologies create a clear, responsive escrow payment platform that can later be connected to backend services, payment rails, logistics APIs, authentication, and persistent databases.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
