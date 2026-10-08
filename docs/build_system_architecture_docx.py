from __future__ import annotations

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "SYSTEM_ARCHITECTURE_FRONTEND_BACKEND.md"
OUTPUT = ROOT / "docs" / "VeriPay_Frontend_Backend_System_Architecture.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    tc_pr.append(shading)


def set_cell_borders(cell, color: str = "D9D9D9") -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.first_child_found_in("w:tcBorders")
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)

    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_font(run, size: int | float | None = None, bold: bool | None = None, color: str | None = None, font: str = "Aptos") -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if color is not None:
        run.font.color.rgb = RGBColor.from_string(color)


def configure_styles(document: Document) -> None:
    styles = document.styles

    normal = styles["Normal"]
    normal.font.name = "Aptos"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
    normal.font.size = Pt(10.5)
    normal.font.color.rgb = RGBColor(0, 0, 0)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.08

    title = styles["Title"]
    title.font.name = "Aptos Display"
    title._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos Display")
    title.font.size = Pt(24)
    title.font.bold = True
    title.font.color.rgb = RGBColor(0, 0, 0)
    title.paragraph_format.space_after = Pt(6)

    for name, size in (("Heading 1", 15), ("Heading 2", 12.5), ("Heading 3", 11.2)):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = styles[name]
        style.font.name = "Aptos"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Aptos")
        style.font.size = Pt(10.3)
        style.paragraph_format.space_after = Pt(3)


def add_header_footer(document: Document) -> None:
    section = document.sections[0]
    header = section.header.paragraphs[0]
    header.text = "VeriPay System Architecture"
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_font(header.runs[0], size=8.5, color="666666")

    footer = section.footer.paragraphs[0]
    footer.text = "Frontend and backend technical documentation"
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_font(footer.runs[0], size=8.5, color="666666")


def add_code_block(document: Document, lines: list[str]) -> None:
    for line in lines:
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.left_indent = Inches(0.22)
        paragraph.paragraph_format.right_indent = Inches(0.12)
        paragraph.paragraph_format.space_after = Pt(0)
        paragraph.paragraph_format.line_spacing = 1.0
        run = paragraph.add_run(line if line else " ")
        set_font(run, size=8.7, font="Consolas", color="111827")
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), "F3F4F6")
        paragraph._p.get_or_add_pPr().append(shading)

    spacer = document.add_paragraph()
    spacer.paragraph_format.space_after = Pt(4)


def add_key_value_table(document: Document, rows: list[tuple[str, str]]) -> None:
    table = document.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    table.autofit = False
    table.columns[0].width = Inches(2.0)
    table.columns[1].width = Inches(4.6)

    hdr = table.rows[0].cells
    hdr[0].text = "Area"
    hdr[1].text = "Description"
    for cell in hdr:
        set_cell_shading(cell, "1F2937")
        set_cell_borders(cell)
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        for paragraph in cell.paragraphs:
            for run in paragraph.runs:
                set_font(run, size=9.5, bold=True, color="FFFFFF")
    set_repeat_table_header(table.rows[0])

    for idx, (left, right) in enumerate(rows):
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
        for cell in cells:
            set_cell_borders(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if idx % 2 == 1:
                set_cell_shading(cell, "F8FAFC")
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(0)
                for run in paragraph.runs:
                    set_font(run, size=9.2)
    document.add_paragraph()


def clean_inline_code(text: str) -> str:
    return text.replace("`", "")


def add_markdown_paragraph(document: Document, text: str, style: str | None = None) -> None:
    text = clean_inline_code(text)
    paragraph = document.add_paragraph(style=style)
    for idx, part in enumerate(re.split(r"(\*\*[^*]+\*\*)", text)):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_font(run, bold=True)
        else:
            run = paragraph.add_run(part)
            set_font(run)


def build_document() -> None:
    markdown = SOURCE.read_text(encoding="utf-8")
    document = Document()

    section = document.sections[0]
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.72)
    section.left_margin = Inches(0.78)
    section.right_margin = Inches(0.78)

    configure_styles(document)
    add_header_footer(document)

    title = document.add_paragraph(style="Title")
    title.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = title.add_run("VeriPay Frontend and Backend System Architecture")
    set_font(run, size=24, bold=True, font="Aptos Display")

    subtitle = document.add_paragraph()
    subtitle.paragraph_format.space_after = Pt(12)
    run = subtitle.add_run("Detailed technical explanation of the current frontend, backend, login, API, data, and runtime system as implemented on 2026-10-08.")
    set_font(run, size=10.5, color="4B5563")

    add_key_value_table(
        document,
        [
            ("Frontend", "React, Vite, TypeScript, Tailwind CSS, and reusable dashboard components."),
            ("Backend", "Node.js HTTP server with signed token authentication and protected API routes."),
            ("Persistence", "Local JSON datastore at data/veripay-db.json for MVP development."),
            ("Primary command", "npm run dev starts both the API server and the Vite frontend."),
        ],
    )

    lines = markdown.splitlines()
    index = 0
    in_code = False
    code_lines: list[str] = []
    skip_title = True

    while index < len(lines):
        raw = lines[index]
        line = raw.rstrip()

        if line.startswith("```"):
            if in_code:
                add_code_block(document, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            index += 1
            continue

        if in_code:
            code_lines.append(line)
            index += 1
            continue

        if not line.strip():
            index += 1
            continue

        if skip_title and line.startswith("# "):
            skip_title = False
            index += 1
            continue

        if line.startswith("# "):
            add_markdown_paragraph(document, line[2:].strip(), "Heading 1")
        elif line.startswith("## "):
            add_markdown_paragraph(document, line[3:].strip(), "Heading 1")
        elif line.startswith("### "):
            add_markdown_paragraph(document, line[4:].strip(), "Heading 2")
        elif re.match(r"^\d+\.\s+", line):
            add_markdown_paragraph(document, re.sub(r"^\d+\.\s+", "", line), "List Number")
        elif line.startswith("- "):
            add_markdown_paragraph(document, line[2:].strip(), "List Bullet")
        else:
            add_markdown_paragraph(document, line, None)

        index += 1

    document.save(OUTPUT)


if __name__ == "__main__":
    build_document()
    print(OUTPUT)
