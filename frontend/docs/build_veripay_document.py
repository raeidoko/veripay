from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "docs" / "VeriPay_Site_Overview.docx"
SHOT_DIR = ROOT / "docs" / "screenshots_clean"

BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
NAVY = RGBColor(15, 23, 42)
MUTED = RGBColor(100, 116, 139)
LIGHT_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "D9E2EC"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
      shd = OxmlElement("w:shd")
      tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_widths(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")
    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    grid = table._tbl.tblGrid
    if grid is None:
        grid = OxmlElement("w:tblGrid")
        table._tbl.insert(0, grid)
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = Inches(widths[idx] / 1440)
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(widths[idx]))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_run_font(run, size=None, color=None, bold=None, italic=None):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def style_doc(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    for side in ("top_margin", "right_margin", "bottom_margin", "left_margin"):
        setattr(section, side, Inches(1))
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.line_spacing = 1.10

    header = section.header.paragraphs[0]
    header.text = ""
    r = header.add_run("VeriPay Site Overview")
    set_run_font(r, size=9, color=MUTED, bold=True)
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT

    footer = section.footer.paragraphs[0]
    footer.text = ""
    r = footer.add_run("Prepared documentation for VeriPay secure escrow payment platform")
    set_run_font(r, size=8.5, color=MUTED)
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def add_para(doc, text="", size=11, bold=False, italic=False, color=None, align=None, after=6, before=0):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.10
    if align is not None:
        p.alignment = align
    run = p.add_run(text)
    set_run_font(run, size=size, color=color, bold=bold, italic=italic)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_number(doc, text):
    p = doc.add_paragraph(style="List Number")
    p.paragraph_format.space_after = Pt(8)
    p.paragraph_format.line_spacing = 1.167
    run = p.add_run(text)
    set_run_font(run, size=11)
    return p


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    set_table_widths(table, [9360])
    cell = table.cell(0, 0)
    set_cell_shading(cell, CALLOUT_FILL)
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(3)
    r = p.add_run(title)
    set_run_font(r, size=11, color=NAVY, bold=True)
    p2 = cell.add_paragraph()
    p2.paragraph_format.space_after = Pt(2)
    p2.paragraph_format.line_spacing = 1.10
    r2 = p2.add_run(body)
    set_run_font(r2, size=10.5)
    add_para(doc, "", after=4)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_widths(table, widths)
    for i, h in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_FILL)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        r = p.add_run(h)
        set_run_font(r, size=10, color=NAVY, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.10
            r = p.add_run(value)
            set_run_font(r, size=10)
    set_table_widths(table, widths)
    add_para(doc, "", after=4)
    return table


def add_caption(doc, label):
    p = add_para(doc, label, size=9.5, italic=True, color=MUTED, after=8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    return p


def add_screenshot_section(doc, title, screenshot_name, caption, paragraphs, bullets=None):
    doc.add_page_break()
    doc.add_heading(title, level=1)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(str(SHOT_DIR / screenshot_name), width=Inches(6.35))
    add_caption(doc, caption)
    for text in paragraphs:
        add_para(doc, text)
    if bullets:
        for text in bullets:
            add_bullet(doc, text)


def build():
    doc = Document()
    style_doc(doc)

    add_para(doc, "VERIPAY PRODUCT OVERVIEW", size=10, bold=True, color=BLUE, after=2)
    add_para(doc, "VeriPay Site Walkthrough and Feature Documentation", size=24, bold=True, color=NAVY, after=4)
    add_para(
        doc,
        "A detailed explanation of the secure escrow payment platform, including screenshots of the core product pages and the user journeys they support.",
        size=13,
        color=MUTED,
        after=16,
    )
    metadata = [
        ("Project", "VeriPay"),
        ("Document type", "Product overview and page walkthrough"),
        ("Prepared for", "Internal review, stakeholder explanation, and product presentation"),
        ("Prepared on", "July 3, 2026"),
        ("Site URL used for screenshots", "http://localhost:3000"),
    ]
    for label, value in metadata:
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(f"{label}: ")
        set_run_font(r, size=11, bold=True)
        r2 = p.add_run(value)
        set_run_font(r2, size=11)

    add_callout(
        doc,
        "Purpose of this document",
        "This document explains what VeriPay does, how the site is organized, and how each screen contributes to the escrow workflow. It is written for someone who has not yet used the product and needs a clear, practical understanding of the experience.",
    )

    doc.add_heading("Executive Summary", level=1)
    add_para(
        doc,
        "VeriPay is a secure escrow payment platform for social commerce transactions. It is designed around a common online-trade problem: a buyer does not want to release money before receiving the item, while a seller does not want to ship before knowing the buyer can pay. VeriPay sits between both parties by holding funds, tracking delivery progress, and providing a structured path for dispute resolution when something goes wrong.",
    )
    add_para(
        doc,
        "The site presents this flow as an operational product dashboard. It includes a transaction pipeline, a seller directory, transaction history, notifications, a dispute center, and digital receipt generation. Together, these pages show how a protected marketplace payment can move from agreement to payment, shipment, delivery, confirmation, payout, or refund.",
    )

    doc.add_heading("Who the Site Serves", level=2)
    for item in [
        "Buyers who need confidence that funds will not be released until an item has been delivered and inspected.",
        "Sellers who need proof that money has been secured before they dispatch goods.",
        "Logistics partners who update shipment milestones and help validate delivery status.",
        "Support or arbitration teams who need evidence, conversation history, and a final decision path for disputes.",
        "Product stakeholders who need a clear view of the escrow workflow, transaction states, trust controls, and customer-facing experience.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Core Product Areas", level=2)
    add_table(
        doc,
        ["Area", "What it does", "Why it matters"],
        [
            ("Active escrow pipeline", "Shows the selected transaction status, current milestone, and settlement state.", "Keeps the user focused on the current stage of the protected order."),
            ("History and receipts", "Lists previous and active transactions, values, statuses, and receipt actions.", "Creates traceability and makes completed orders auditable."),
            ("Verified sellers", "Displays merchants with trust scores, sales history, categories, and quick start actions.", "Helps buyers begin safer transactions from a vetted seller context."),
            ("Dispute center", "Collects dispute claims, evidence, mediated chat, and arbitration decisions.", "Protects both sides when delivery or product condition is contested."),
            ("Digital receipt", "Summarizes the settled transaction, parties, item terms, fees, and security reference.", "Provides proof of settlement after funds have been released."),
        ],
        [1700, 3700, 3960],
    )

    add_screenshot_section(
        doc,
        "Page 1: Active Escrow Pipeline",
        "01-active-escrow-pipeline.png",
        "Figure 1. Active escrow pipeline showing the selected transaction, status timeline, settlement state, and alert feed.",
        [
            "The Active Escrow Pipeline is the operational center of VeriPay. It shows the currently selected transaction and explains what stage the order is in. The page combines navigation tabs, the escrow workflow card, transaction state details, and notifications.",
            "In the screenshot, the selected transaction is in a disputed state. The pipeline has moved through the major milestones and the status badge shows that the order is no longer a normal delivery flow. This is important because escrow systems must be explicit about whether money is still protected, released, refunded, or under review.",
            "The page also makes the account state easy to understand. A user can immediately see whether the transaction is pending payment, secured, in transit, delivered, completed, refunded, or disputed.",
        ],
        [
            "The timeline communicates the transaction's progress from draft through settlement.",
            "The status badge gives a quick answer about the current risk or completion state.",
            "The alert feed keeps important escrow and dispute updates visible alongside the active order.",
        ],
    )

    add_screenshot_section(
        doc,
        "Page 2: History and Receipts",
        "02-history-receipts.png",
        "Figure 2. Transaction history page showing completed, disputed, in-transit, and pending-payment orders.",
        [
            "The History and Receipts page acts as the transaction ledger. It gives users a simple way to review all orders, inspect their current status, and jump back into an active transaction. Each row includes the order ID, product name, merchant, date, value, status, and available action.",
            "This page is especially useful because escrow products need strong recordkeeping. A buyer may want to confirm which orders are pending, which are complete, and which are disputed. A seller may want to identify which transaction is still waiting for payment before shipping. A support team may also use this view to orient itself before reviewing a case.",
            "Completed transactions show a Receipt button. Active transactions show a Track Order action. This distinction prevents users from treating closed orders like active orders while still preserving access to settlement proof.",
        ],
    )

    add_screenshot_section(
        doc,
        "Page 3: Verified Sellers",
        "03-verified-sellers.png",
        "Figure 3. Verified seller directory with merchant cards, trust scores, categories, and escrow-start actions.",
        [
            "The Verified Sellers page is the entry point for starting a protected order from a merchant profile. It shows vendors, categories, phone numbers, star ratings, successful sales, active escrows, and trust scores. The user can search by merchant name or handle and filter by category.",
            "This screen matters because social commerce usually begins on platforms like Instagram, WhatsApp, TikTok, or direct messages. VeriPay gives that informal discovery process a safer checkout layer. Instead of sending money directly to an unknown vendor, the buyer can start an escrow order from a seller card.",
            "Each seller card includes a Start Escrow Order button and a QR-style quick action. These actions can support seller payment links, merchant QR codes, identity checks, and a verified seller onboarding system.",
        ],
        [
            "Trust scores give buyers an at-a-glance risk signal.",
            "Category filters help users browse merchants by product type.",
            "Active escrow counts show that merchants are participating in the protected flow.",
        ],
    )

    add_screenshot_section(
        doc,
        "Page 4: Dispute Center",
        "04-dispute-center.png",
        "Figure 4. Dispute center showing an active claim, evidence submitted by both parties, mediated chat, and arbitration actions.",
        [
            "The Dispute Center is where VeriPay handles transactions that cannot be settled through normal delivery confirmation. The page shows the disputed order, the reason for the claim, the value under review, evidence submitted by the buyer and seller, and a mediated discussion thread.",
            "In the screenshot, the buyer has filed a claim for a wrong or damaged item. The center panel explains the issue and displays uploaded proof. The seller counter-proof area shows the seller's evidence. This structure helps reduce confusion because each side's materials are separated and labeled.",
            "The arbitration controls at the top right show how an authorized reviewer can resolve the case. The reviewer may refund the buyer or release payment to the seller based on the claim, evidence, merchant response, and delivery record.",
        ],
        [
            "Buyer proof supports refund claims and product-condition complaints.",
            "Seller counter-proof protects merchants from false or incomplete claims.",
            "Mediated discussion keeps case communication tied to the transaction record.",
        ],
    )

    add_screenshot_section(
        doc,
        "Page 5: Digital Receipt",
        "05-digital-receipt-modal.png",
        "Figure 5. Digital receipt modal with transaction value, settlement date, parties, logistics reference, item terms, fees, and verification badge.",
        [
            "The Digital Receipt modal is the proof document generated after a transaction is completed. It summarizes the buyer, seller, item, delivery partner, tracking reference, transaction value, escrow fee, and total settled amount.",
            "The receipt is important because escrow systems do not end at payment release. Users still need a clear post-settlement record that shows what was agreed, who participated, when settlement occurred, and how much was processed. This can support buyer records, seller bookkeeping, customer support, and dispute follow-up.",
            "The modal includes print support, so it can be exported or physically saved. The same receipt structure can support downloadable PDF receipts, signed settlement confirmations, email receipts, or account-level transaction statements.",
        ],
    )

    doc.add_page_break()
    doc.add_heading("End-to-End User Journey", level=1)
    add_para(doc, "A typical VeriPay transaction can be understood as a sequence of protected steps:")
    for step in [
        "The buyer chooses a verified seller or begins a custom escrow draft.",
        "The buyer and seller agree on product details, price, delivery partner, delivery window, and inspection terms.",
        "The buyer deposits the agreed funds into escrow, including any platform fee.",
        "The seller receives confirmation that funds are secured and ships the product.",
        "The logistics partner updates the delivery status as the package moves through the route.",
        "The buyer receives the item and either confirms delivery or opens a dispute.",
        "If confirmed, VeriPay releases funds to the seller and generates a receipt.",
        "If disputed, VeriPay locks the funds, gathers evidence, and waits for an arbitration decision.",
    ]:
        add_number(doc, step)

    doc.add_heading("Operational Platform Capabilities", level=1)
    add_para(
        doc,
        "VeriPay presents the main operational capabilities required for a protected social commerce payment experience. The interface is organized around transaction state, seller trust, delivery visibility, dispute evidence, settlement decisions, and receipts.",
    )
    add_table(
        doc,
        ["Capability", "What the platform provides", "Operational value"],
        [
            ("Escrow funding", "Records a protected payment state before the seller releases goods.", "Reduces payment risk for sellers and delivery risk for buyers."),
            ("Delivery tracking", "Displays dispatch, transit, delivery, and settlement milestones.", "Keeps both parties aligned on the current order stage."),
            ("Disputes", "Organizes claims, evidence, merchant responses, and reviewer decisions.", "Creates a fair case-management path when the order is contested."),
            ("Receipts", "Provides a printable settlement record with item, parties, fees, and references.", "Supports customer support, bookkeeping, and post-order proof."),
            ("Seller trust", "Displays trust signals, sales history, ratings, and active escrow counts.", "Helps buyers choose merchants with clearer confidence signals."),
        ],
        [2200, 3300, 3860],
    )

    doc.add_heading("Key Value Proposition", level=1)
    add_para(
        doc,
        "The core value of VeriPay is trust. The product makes social commerce feel less risky by replacing direct, informal payment with a structured agreement, a payment hold, delivery milestones, and a review path. Buyers get confidence that their money is protected. Sellers get confidence that payment exists before shipping. Support teams get a documented record when a case needs review.",
    )
    add_callout(
        doc,
        "Practical takeaway",
        "VeriPay is strongest when presented as a trust layer for informal digital trade: it does not replace social selling, messaging, or logistics, but it adds the protected payment and accountability layer those channels often lack.",
    )

    doc.add_heading("Recommended Next Improvements", level=1)
    for item in [
        "Add authentication and role-based permissions so buyer, seller, logistics, and arbitrator actions are separated securely.",
        "Connect escrow funding to a real payment provider or ledger service with server-side transaction validation.",
        "Store transactions, notifications, disputes, and receipts in a backend database instead of browser localStorage.",
        "Add upload support for dispute evidence, including image, document, and video files.",
        "Generate downloadable PDF receipts and case summaries for completed and disputed transactions.",
        "Add mobile-specific QA because the target users will likely start many transactions from phones.",
    ]:
        add_bullet(doc, item)

    doc.add_heading("Conclusion", level=1)
    add_para(
        doc,
        "VeriPay already communicates the main mechanics of a secure escrow platform. The dashboard explains the transaction state, the seller directory explains how protected orders begin, the history page preserves the order record, the dispute center handles exceptions, and the receipt modal closes the loop with proof of settlement. Together, the screens form a coherent product story: safer trade through staged payment, transparent delivery tracking, and structured resolution.",
    )

    doc.save(OUT)


if __name__ == "__main__":
    build()
